#!/usr/bin/env python3
"""Synthetic, offline regressions; NO test evidence represents a real stock search."""
from __future__ import annotations
import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'shared/scripts'))
from editorial_contract import apply,audit as audit_notes
from audit_stock_execution import audit as audit_stock

DEMO_URL='https://www.pexels.com/video/12345/'  # structural test only; not a verified live asset


def blueprint(kind='video',stock=True):
    b={'id':'R1-B1','title':'Conoce una tarea','format':'talento + apoyo visual',
       'direction':'Mostrar el apoyo visual en el momento en que comienza la demostracion.',
       'stock_decision':{'needed':stock,'reason':'La accion central requiere una escena especifica para comprender el ejemplo.' if stock else 'El diagrama demuestra el proceso con mayor claridad que una foto.'},
       'visuals':[{'kind':'stock_video' if stock else 'diagram', 'purpose':'Reconocer la accion concreta que se describe en el ejemplo.'}],
       'stock_assets':['AS1'] if stock else []}
    if kind=='video' and stock:
        b['stock_video_links']=[{'asset_id':'AS1','title':'Clip elegido','url':DEMO_URL}]
        b['timecode']='00:00-00:20'
    return {'resource_id':'R1','resource_type':kind,'blocks':[b]}


def evidence(kind='stock_video',status='selected'):
    r={'asset_id':'AS1','block_id':'R1-B1','kind':kind,'provider':'Proveedor simulado',
       'query':'persona trabajando en oficina con documentos','searched_at':'2026-10-26T13:00:00+01:00',
       'status':status,'evidence':{'tool':'TEST_ONLY','result_ref':'SYNTHETIC_12345'},
       'candidates':[{'title':'Clip elegido','url':DEMO_URL}],
       'selected':{'title':'Clip elegido','url':DEMO_URL}}
    if status=='no_match':
        r['attempted_queries']=['persona trabajando en oficina','colaborador revisando documentos']
        r.pop('selected');r['candidates']=[]
    if status=='access_blocked':
        r.pop('selected');r['candidates']=[];r['error']='Proveedor no conectado en fixture de test'
    return {'resource_id':'R1','searches':[r]}


def stub_docx(path,kind='video'):
    d=Document()
    meta=d.add_table(rows=4,cols=2)
    for i,(k,v) in enumerate([('Tipo de Recurso',kind),('Duraci\u00f3n','2 min'),('Estilo/Edici\u00f3n','Estilo claro'),('Descripci\u00f3n/Objetivo','Practicar un paso')]):
        meta.rows[i].cells[0].text=k;meta.rows[i].cells[1].text=v
    table=d.add_table(rows=2,cols=2)
    table.rows[0].cells[0].text='Notas de Edici\u00f3n'
    table.rows[0].cells[1].text='Contenido del Recurso'
    table.rows[1].cells[0].text='Por completar'
    table.rows[1].cells[1].text='Este texto forma parte del recurso y se conserva.'
    d.save(path)


class Guardrails(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.root=Path(self.tmp.name)
        self.src=self.root/'input.docx';self.doc=self.root/'out.docx'
        stub_docx(self.src)
    def tearDown(self):self.tmp.cleanup()

    def test_valid_stock_video_notes_and_evidence(self):
        c=blueprint();apply(self.src,c,self.doc)
        self.assertTrue(audit_notes(self.doc,c)['pass_delivery'])
        self.assertEqual(audit_stock(c,evidence(),self.doc)['state'],'READY_FOR_VISUAL_REVIEW')
    def test_notes_not_canonical_fails(self):
        c=blueprint();apply(self.src,c,self.doc)
        changed=Document(self.doc);changed.tables[1].rows[1].cells[0].paragraphs[2].text+=' Ademas, explicar durante seis minutos todos los criterios.'
        changed.save(self.doc)
        self.assertFalse(audit_notes(self.doc,c)['pass_delivery'])
    def test_stock_missing_search_fails(self):
        c=blueprint();apply(self.src,c,self.doc)
        report=audit_stock(c,{'searches':[]},self.doc)
        self.assertEqual(report['state'],'INVALID')
        self.assertIn('STOCK_SEARCH_NOT_EXECUTED',[r['code'] for r in report['errors']])
    def test_access_blocked_not_ready(self):
        c=blueprint();c['blocks'][0].pop('stock_video_links')
        apply(self.src,c,self.doc)
        self.assertEqual(audit_stock(c,evidence(status='access_blocked'),self.doc)['state'],'IMPLEMENTATION_PENDING')
    def test_no_match_not_ready(self):
        c=blueprint();c['blocks'][0].pop('stock_video_links')
        apply(self.src,c,self.doc)
        self.assertEqual(audit_stock(c,evidence(status='no_match'),self.doc)['state'],'IMPLEMENTATION_PENDING')
    def test_no_match_without_retry_fails(self):
        c=blueprint();apply(self.src,c,self.doc)
        log=evidence(status='no_match');log['searches'][0]['attempted_queries']=['solo consulta']
        self.assertEqual(audit_stock(c,log,self.doc)['state'],'INVALID')
    def test_non_stock_with_good_reason_passes(self):
        c=blueprint(stock=False);apply(self.src,c,self.doc)
        self.assertEqual(audit_stock(c,{'searches':[]},self.doc)['state'],'READY_FOR_VISUAL_REVIEW')
    def test_image_chosen_but_not_embedded_fails(self):
        c=blueprint(kind='rise',stock=True);c['blocks'][0].pop('stock_video_links',None)
        c['blocks'][0]['visuals'][0]['kind']='stock_image'
        apply(self.src,c,self.doc)
        log=evidence(kind='stock_image')
        report=audit_stock(c,log,self.doc)
        codes={x['code'] for x in report['errors']}
        self.assertIn('STOCK_IMAGE_NOT_EMBEDDED_IN_BLOCK',codes)
        self.assertIn('STOCK_IMAGE_LINK_MISSING_IN_CONTENT',codes)
    def test_tabs_missing_panels_fails(self):
        c=blueprint(stock=False);c['blocks'][0]['block_kind']='tabs'
        apply(self.src,c,self.doc)
        self.assertIn('PANEL_DECISIONS_MISSING',[x['code'] for x in audit_stock(c,{'searches':[]},self.doc)['errors']])
    def test_stock_without_reason_fails(self):
        c=blueprint();c['blocks'][0]['stock_decision']['reason']='No aplica'
        with self.assertRaises(ValueError):apply(self.src,c,self.doc)
    def test_video_link_not_from_search_fails(self):
        c=blueprint();c['blocks'][0]['stock_video_links'][0]['url']='https://www.pexels.com/video/99999/'
        apply(self.src,c,self.doc)
        self.assertIn('VIDEO_LINK_NOT_FROM_SEARCH',[x['code'] for x in audit_stock(c,evidence(),self.doc)['errors']])
    def test_all_no_stock_three_blocks_requires_second_review(self):
        c=blueprint(stock=False)
        for idx in range(2,4):
            b=copy.deepcopy(c['blocks'][0]);b['id']='R1-B'+str(idx)
            b['title']='Explica el paso '+str(idx)
            c['blocks'].append(b)
        report=audit_stock(c,{'searches':[]})
        self.assertEqual(report['state'],'IMPLEMENTATION_PENDING')
        log={'searches':[], 'all_no_stock_review':{'reviewer':'visual-stock-reviewer','accepted':True,'rationale':'Los diagramas ilustran tres relaciones conceptuales y no requieren escenas fotograficas.'}}
        self.assertEqual(audit_stock(c,log)['state'],'READY_FOR_VISUAL_REVIEW')

    def test_notes_over_60_words_fails_contract(self):
        c=blueprint(stock=False)
        c['blocks'][0]['direction']='Mostrar el diagrama cuando se explique el primer paso. '*16
        with self.assertRaises(ValueError):apply(self.src,c,self.doc)

if __name__=='__main__':unittest.main()
