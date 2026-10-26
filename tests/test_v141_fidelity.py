#!/usr/bin/env python3
"""Offline fidelity regressions. Synthetic sources do NOT certify real course fidelity."""
import copy
import hashlib
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from docx import Document

SUITE=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(SUITE/'shared/scripts'))
from fidelity_gate import verify,prepare,digest,word_content_units,source_segments

SOURCE_TEXT=[
    'Para solicitar la licencia se requiere aprobacion del responsable, completar una capacitacion y generar un ticket.',
    'El portal de ayuda esta disponible desde la intranet corporativa.',
]


def write_json(p,data):
    p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


class FidelityRegressions(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.root=Path(self.tmp.name)
        for n in ('qa','sources','outputs'):(self.root/n).mkdir()
        self.source=self.root/'sources/material.txt'
        self.source.write_text('\n'.join(SOURCE_TEXT)+'\n',encoding='utf-8')
        self.word=self.root/'outputs/R1.docx'
        self.make_docx()
        self.registry=self.root/'qa/source-catalog.json'
        self.coverage=self.root/'qa/course-coverage-contract.json'
        self.ledger=self.root/'qa/R1.fidelity-ledger.json'
        self.review=self.root/'qa/R1.fidelity-review.json'
        write_json(self.registry,{'schema_version':'1.4.1','sources':[
            {'id':'F1','path':'sources/material.txt','factual':True,'role':'principal','priority':1,'sha256':digest(self.source)}]})
        write_json(self.coverage,{'schema_version':'1.4.1','topics':[
            {'id':'T1','source_id':'F1','locator':'line:1','quote':SOURCE_TEXT[0],
             'usage':'content','required':True,'resource_ids':['R1']},
            {'id':'T2','source_id':'F1','locator':'line:2','quote':SOURCE_TEXT[1],
             'usage':'content','required':True,'resource_ids':['R1']}]})
        prepare(self.root,'R1',self.word,self.registry,self.ledger)
        self.build_ledger_review()
    def tearDown(self):self.tmp.cleanup()
    def make_docx(self):
        doc=Document()
        metadata=doc.add_table(rows=2,cols=2)
        metadata.cell(0,0).text='Tipo de Recurso'
        metadata.cell(0,1).text='Video explicativo'
        metadata.cell(1,0).text='Objetivo'
        metadata.cell(1,1).text='Relacionar pasos'
        tb=doc.add_table(rows=3,cols=2)
        tb.rows[0].cells[0].text='Notas de Edicion'
        tb.rows[0].cells[1].text='Contenido del Recurso'
        tb.rows[1].cells[0].text='1. Tu primera licencia\nVideo: talento.\nMostrar los tres pasos.'
        tb.rows[1].cells[1].text='Solicita aprobacion de tu responsable, completa la capacitacion y genera un ticket.'
        tb.rows[2].cells[0].text='2. Ayuda disponible\nVideo: pantalla.\nMostrar el acceso a intranet.'
        tb.rows[2].cells[1].text='Consulta los materiales desde la intranet corporativa.'
        doc.save(self.word)
    def build_ledger_review(self):
        ledger=json.loads(self.ledger.read_text())
        for n,u in enumerate(ledger['units'],1):
            u['classification']='factual'
            u['claim_ids']=['C'+str(n)]
        ledger['claims']=[
            {'id':'C1','unit_id':'R1-U001','text':ledger['units'][0]['text'],
             'verdict':'RESPALDADA','evidence':[{'source_id':'F1','locator':'line:1',
              'quote':SOURCE_TEXT[0], 'rationale':'La fuente exige exactamente tres acciones previas y el recurso las mantiene.'}]},
            {'id':'C2','unit_id':'R1-U002','text':ledger['units'][1]['text'],
             'verdict':'RESPALDADA','evidence':[{'source_id':'F1','locator':'line:2',
              'quote':SOURCE_TEXT[1], 'rationale':'La fuente indica que el portal esta en la intranet, sin inventar enlace.'}]},
        ]
        ledger['coverage_complete']=True
        ledger['topic_coverage']=[
            {'topic_id':'T1','claim_ids':['C1'],'rationale':'La unidad uno cubre los tres pasos exigidos por la fuente.'},
            {'topic_id':'T2','claim_ids':['C2'],'rationale':'La unidad dos mantiene el canal de acceso indicado.'},
        ]
        write_json(self.ledger,ledger)
        report={
            'agent':'source-fidelity-reviewer','resource_id':'R1','resource':'R1','status':'VERDE',
            'decision':'APROBADO','review_complete':True,'word_sha256':digest(self.word),
            'source_hashes':{'F1':digest(self.source)},'reviewed_at':'2026-10-26T13:00:00+01:00',
            'reviewed_unit_ids':['R1-U001','R1-U002'],'reviewed_claim_ids':['C1','C2'],
            'reviewed_topic_ids':['T1','T2'],
            'claim_assessments':[
                {'claim_id':'C1','verdict':'RESPALDADA','explanation':'Compare los tres requisitos uno a uno con el material base; ninguno falta o cambia.'},
                {'claim_id':'C2','verdict':'RESPALDADA','explanation':'La fuente nombra la intranet como canal; el texto no promete URL o tiempo.'},
            ],
            'topic_assessments':[
                {'topic_id':'T1','covered':True,'explanation':'Se incluyen aprobacion, capacitacion y ticket en ese orden.'},
                {'topic_id':'T2','covered':True,'explanation':'Se recupera la intranet sin introducir condiciones nuevas.'},
            ],
            'semantic_checks':{x:True for x in ('claim_completeness','evidence_support','conditions_retained',
                              'source_precedence_reviewed','conflicts_resolved','no_outside_facts')},
            'review_notes':'Se revisaron las dos afirmaciones del recurso, las condiciones de la licencia y el acceso. La fuente tiene prioridad unica y los ejemplos no anaden hechos externos.',
            'issues':[],
        }
        write_json(self.review,report)
    def audit(self):
        return verify(self.root,'R1',self.word,self.registry,self.coverage,self.ledger,self.review)
    def codes(self):return ' | '.join(self.audit()['errors'])
    def rewrite(self,path,func):
        d=json.loads(path.read_text());func(d);write_json(path,d)

    def test_valid_document_claims_approved(self):
        r=self.audit();self.assertTrue(r['pass_machine_gates'],r['errors'])
        self.assertEqual(r['machine_checks']['required_topics'],2)
    def test_word_changed_rejects_old_review(self):
        doc=Document(self.word);doc.tables[-1].rows[1].cells[1].text+=' Sin ninguna espera.';doc.save(self.word)
        self.assertIn('FIDELITY_STALE_WORD',self.codes())
    def test_source_changed_rejects_old_review(self):
        self.source.write_text(self.source.read_text()+'\nLa activacion es inmediata.\n')
        self.assertIn('FIDELITY_STALE_SOURCES',self.codes())
    def test_fake_quote_rejected(self):
        self.rewrite(self.ledger,lambda d:d['claims'][0]['evidence'][0].update({'quote':'La activacion tarda exactamente 24 horas.'}))
        self.assertIn('el fragmento NO aparece',self.codes())
    def test_wrong_location_rejected(self):
        self.rewrite(self.ledger,lambda d:d['claims'][0]['evidence'][0].update({'locator':'line:2'}))
        self.assertIn('el fragmento NO aparece',self.codes())
    def test_claim_without_evidence_rejected(self):
        self.rewrite(self.ledger,lambda d:d['claims'][0].update({'evidence':[]}))
        self.assertIn('CLAIM_NO_EVIDENCE',self.codes())
    def test_altered_fact_rejected(self):
        self.rewrite(self.ledger,lambda d:d['claims'][0].update({'verdict':'ALTERADA'}))
        self.assertIn('CLAIM_NOT_SUPPORTED',self.codes())
    def test_missing_claim_rejected(self):
        self.rewrite(self.ledger,lambda d:d['units'][0].update({'claim_ids':[]}))
        self.assertIn('UNIT_NO_CLAIMS',self.codes())
    def test_missing_required_topic_rejected(self):
        self.rewrite(self.ledger,lambda d:d['topic_coverage'].pop())
        self.assertIn('OMITTED_SOURCE_TOPICS',self.codes())
    def test_source_guideline_not_allowed_as_evidence(self):
        self.rewrite(self.registry,lambda d:d['sources'][0].update({'role':'guideline','factual':False}))
        self.assertIn('no fuente factual',self.codes())
    def test_missing_independent_review_rejected(self):
        self.review.unlink()
        self.assertIn('INDEPENDENT_FIDELITY_REVIEW_MISSING',self.codes())
    def test_generic_green_without_claim_assessments_rejected(self):
        self.rewrite(self.review,lambda d:d.update({'claim_assessments':[]}))
        self.assertIn('INDEPENDENT_CLAIM_ASSESSMENTS_INCOMPLETE',self.codes())
    def test_unreviewed_topic_rejected(self):
        self.rewrite(self.review,lambda d:d['topic_assessments'].pop())
        self.assertIn('INDEPENDENT_TOPIC_ASSESSMENTS_INCOMPLETE',self.codes())
    def test_no_fact_source_requires_reason_and_review(self):
        self.rewrite(self.ledger,lambda d:d['units'][0].update({'classification':'nonfactual','claim_ids':[], 'nonfactual_reason':'No aplica'}))
        self.assertIn('UNIT_NO_REASON',self.codes())
    def test_source_excerpt_is_not_only_a_source_id(self):
        self.rewrite(self.ledger,lambda d:d['claims'][0]['evidence'][0].update({'quote':'F1'}))
        self.assertIn('falta fragmento original',self.codes())
    def test_gate_cli_and_course_coverage(self):
        index={'schema_version':'1.4.1','resources':[{'id':'R1','docx':'outputs/R1.docx'}]}
        write_json(self.root/'qa/course-resource-index.json',index)
        script=SUITE/'shared/scripts/fidelity_gate.py'
        p=subprocess.run([sys.executable,str(script),'audit','--project-dir',str(self.root),
                          '--resource-id','R1','--docx','outputs/R1.docx','--strict',
                          '--report','qa/R1.fidelity-preflight.json'],capture_output=True,text=True)
        self.assertEqual(p.returncode,0,p.stdout+p.stderr)
        self.assertTrue(json.loads((self.root/'qa/R1.fidelity-preflight.json').read_text())['pass_machine_gates'])
        whole=SUITE/'shared/scripts/audit_course_coverage.py'
        p=subprocess.run([sys.executable,str(whole),'--project-dir',str(self.root),'--strict'],capture_output=True,text=True)
        self.assertEqual(p.returncode,0,p.stdout+p.stderr)
        self.rewrite(self.ledger,lambda d:d['topic_coverage'].pop())
        p=subprocess.run([sys.executable,str(whole),'--project-dir',str(self.root),'--strict'],capture_output=True,text=True)
        self.assertNotEqual(p.returncode,0)
    def test_final_qa_does_not_trust_generic_green_without_report(self):
        # Ignore Content QA's real semantic relevance: the test checks only the required source gate.
        from_path=SUITE/'shared/references/content-qa-control-register.json'
        reg=json.loads(from_path.read_text())
        reviewer=json.loads(self.review.read_text())
        reviewer['content_qa_checks']=[{'id':c['id'],'result':'NO_APLICA',
                                     'evidence':'Prueba sintentica de agregador; no representa revision QA real.',
                                     'slides':c['source_slides']} for c in reg['controls']]
        write_json(self.review,reviewer)
        agg=SUITE/'skills/final-qa/scripts/aggregate_reports.py'
        cmd=[sys.executable,str(agg),str(self.review),'--project-dir',str(self.root),
             '--resource-id','R1','--docx','outputs/R1.docx',
             '--qa-source-root',str(SUITE/'shared')]
        p=subprocess.run(cmd,capture_output=True,text=True)
        self.assertEqual(p.returncode,0,p.stderr)
        result=json.loads(p.stdout)
        self.assertEqual(result['status'],'VERDE',result)
        self.rewrite(self.review,lambda d:d.update({'claim_assessments':[]}))
        result=json.loads(subprocess.run(cmd,capture_output=True,text=True).stdout)
        self.assertEqual(result['status'],'ROJO')
        self.assertIn('INDEPENDENT_CLAIM_ASSESSMENTS_INCOMPLETE',
                      ' '.join(result['source_fidelity']['errors']))


    def test_unregistered_source_is_blocked(self):
        (self.root/'sources/extra-manual.pdf').write_bytes(b'%PDF not a source yet')
        self.assertIn('UNREGISTERED_SOURCE_FILES',self.codes())

    def test_source_docx_locator(self):
        source=self.root/'sources/evidence.docx'
        d=Document();d.add_paragraph('La aprobacion del responsable es obligatoria.');d.save(source)
        segments=source_segments(source)
        self.assertEqual(segments['paragraph:1'],'La aprobacion del responsable es obligatoria.')
        self.assertNotIn('paragraph:3',segments)

    def test_source_pptx_locator(self):
        from pptx import Presentation
        src=self.root/'sources/evidence.pptx'
        p=Presentation();slide=p.slides.add_slide(p.slide_layouts[5]);slide.shapes.title.text='Proceso aprobado'
        t=slide.shapes.add_textbox(0,0,1000000,300000)
        t.text='La capacitacion es obligatoria para la licencia.'
        p.save(src)
        segments=source_segments(src)
        self.assertIn('Proceso aprobado',segments['slide:1'])
        self.assertIn('La capacitacion es obligatoria',segments['slide:1'])

    def test_missing_course_coverage_contract_is_blocked(self):
        self.coverage.unlink()
        self.assertIn('COURSE_COVERAGE',self.codes())

    def test_noncontent_topic_cannot_justify_fact(self):
        def alter_topic(d):
            d['topics'][0]['usage']='reference'
        self.rewrite(self.coverage,alter_topic)
        self.assertIn('TOPIC_NON_CONTENT_USED_AS_FACT',self.codes())

    def test_review_omitting_one_claim_assessment_is_blocked(self):
        self.rewrite(self.review,lambda d:d['claim_assessments'].pop())
        self.assertIn('INDEPENDENT_CLAIM_ASSESSMENTS_INCOMPLETE',self.codes())

    def test_intake_requires_structured_registry(self):
        p=SUITE/'skills/course-production-orchestrator/scripts/init_course_project.py'
        init=subprocess.run([sys.executable,str(p),'--course-name','Prueba sin fuentes',
                            '--out',str(self.root/'new'), '--resource-count','1'],capture_output=True,text=True)
        self.assertEqual(init.returncode,0,init.stderr)
        q=Path(init.stdout.strip())
        self.assertTrue((q/'qa/source-catalog.json').exists())
        val=SUITE/'skills/course-production-orchestrator/scripts/validate_course_intake.py'
        out=subprocess.run([sys.executable,str(val),str(q)],capture_output=True,text=True)
        self.assertNotEqual(out.returncode,0)
        self.assertIn('CATALOGO DE FUENTES VACIO',out.stdout)

if __name__=='__main__':unittest.main()
